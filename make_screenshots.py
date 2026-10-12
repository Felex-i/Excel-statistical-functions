#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Делает PNG-картинку таблицы-примера с каждого листа excel_statistical_functions.xlsx.

Способ 1: Excel копирует диапазон «как картинку» -> забираем из буфера обмена (настоящий вид Excel).
Способ 2 (если способ 1 не сработал): картинка рисуется самим скриптом
          по значениям, которые посчитал Excel (похоже на Excel, но не пиксель в пиксель).

Нужны: Windows, Microsoft Excel,  pip install pywin32 openpyxl pillow
Запуск:  python make_screenshots.py            (или:  python make_screenshots.py другой_файл.xlsx)
"""
import pathlib
import re
import sys
import time

import openpyxl
from PIL import Image, ImageDraw, ImageFont, ImageGrab

SCALE = 2  # для четкости картинок


# ------------------------------------------------------------- запасная отрисовка
def get_font(size, bold=False):
    names = (["calibrib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold
             else ["calibri.ttf", "arial.ttf", "DejaVuSans.ttf"])
    for n in names:
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            continue
    return ImageFont.load_default()


def wrap(draw, text, font, max_w):
    lines = []
    for para in str(text).split("\n"):
        cur = ""
        for word in para.split(" "):
            trial = (cur + " " + word).strip()
            if draw.textlength(trial, font=font) <= max_w:
                cur = trial
                continue
            if cur:
                lines.append(cur)
            cur = word
            while draw.textlength(cur, font=font) > max_w and len(cur) > 1:  # слишком длинное слово
                k = len(cur)
                while k > 1 and draw.textlength(cur[:k], font=font) > max_w:
                    k -= 1
                lines.append(cur[:k])
                cur = cur[k:]
        lines.append(cur)
    return lines or [""]


def col_letter(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def add_headings(cells, widths_px, first_row, first_col):
    """Добавляет сверху буквы столбцов (A, B, C...), слева номера строк (1, 2, 3...)."""
    gray = (232, 232, 232)

    def lab(t):
        return {"text": t, "bold": False, "fill": gray, "center": True}

    top = [lab("")] + [lab(col_letter(first_col + i)) for i in range(len(widths_px))]
    body = [[lab(str(first_row + r))] + row for r, row in enumerate(cells)]
    return [top] + body, [36] + list(widths_px)


def render_table(cells, widths_px, path):
    """cells[r][c] = dict(text, bold, fill(rgb|None), center?); widths_px - ширины колонок в пикселях."""
    fs = 15 * SCALE
    pad = 5 * SCALE
    tmp = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    widths = [w * SCALE for w in widths_px]
    row_h = []
    wrapped = []
    for row in cells:
        rl, h = [], 20 * SCALE
        for ci, cell in enumerate(row):
            f = get_font(fs, cell["bold"])
            lines = wrap(tmp, cell["text"], f, widths[ci] - 2 * pad)
            rl.append(lines)
            h = max(h, len(lines) * (fs + 4 * SCALE) + 2 * pad)
        wrapped.append(rl)
        row_h.append(h)
    W, H = sum(widths) + 1, sum(row_h) + 1
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    y = 0
    for ri, row in enumerate(cells):
        x = 0
        for ci, cell in enumerate(row):
            if cell["fill"]:
                d.rectangle([x, y, x + widths[ci], y + row_h[ri]], fill=cell["fill"])
            d.rectangle([x, y, x + widths[ci], y + row_h[ri]], outline=(208, 208, 208))
            f = get_font(fs, cell["bold"])
            ty = y + pad
            for line in wrapped[ri][ci]:
                tx = x + pad
                if cell.get("center"):
                    tx = x + (widths[ci] - d.textlength(line, font=f)) / 2
                d.text((tx, ty), line, fill="black", font=f)
                ty += fs + 4 * SCALE
            x += widths[ci]
        y += row_h[ri]
    img.save(path)


# ------------------------------------------------------------------- работа с Excel
def is_blank(img):
    return img.convert("RGB").getextrema() == ((255, 255), (255, 255), (255, 255))


def capture_with_excel(ws, rng, png):
    """Возвращает None при успехе или текст причины неудачи."""
    try:  # чтобы в картинку попали буквы столбцов и номера строк
        ws.PageSetup.PrintHeadings = True
        ws.PageSetup.PrintGridlines = True
    except Exception:
        pass
    try:
        rng.CopyPicture(Appearance=2, Format=2)  # 2 = «как при печати» (учитывает заголовки)
    except Exception as e:  # noqa
        return f"CopyPicture: {e}"
    time.sleep(0.5)
    img = ImageGrab.grabclipboard()
    if not isinstance(img, Image.Image):
        return "буфер обмена пуст"
    if is_blank(img):
        return "получилась пустая (белая) картинка"
    img.save(png)
    return None


def capture_by_drawing(ws, rng, png):
    r0, c0 = rng.Row, rng.Column
    nr, nc = rng.Rows.Count, rng.Columns.Count
    widths = [int(ws.Columns(c0 + i).ColumnWidth * 7 + 5) for i in range(nc)]
    cells = []
    for r in range(nr):
        row = []
        for c in range(nc):
            cell = ws.Cells(r0 + r, c0 + c)
            fill = None
            try:
                if cell.Interior.ColorIndex != -4142:  # есть заливка
                    v = int(cell.Interior.Color)
                    fill = (v & 255, (v >> 8) & 255, (v >> 16) & 255)  # BGR -> RGB
            except Exception:
                pass
            row.append({"text": cell.Text, "bold": bool(cell.Font.Bold), "fill": fill})
        cells.append(row)
    cells, widths = add_headings(cells, widths, r0, c0)
    render_table(cells, widths, png)


def main():
    import win32com.client as win32

    xlsx = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "excel_statistical_functions.xlsx").resolve()
    out = pathlib.Path("screenshots").resolve()
    out.mkdir(exist_ok=True)

    meta = openpyxl.load_workbook(xlsx, read_only=True)["Все функции"]
    items = [(r[0], r[1], r[8], r[9]) for r in meta.iter_rows(min_row=2, values_only=True)
             if r[8] and r[9]]

    excel = win32.Dispatch("Excel.Application")
    excel.Visible = True
    excel.DisplayAlerts = False
    wb = excel.Workbooks.Open(str(xlsx))
    excel.CalculateFull()

    n_excel = n_draw = n_err = 0
    for num, name, sheet, addr in items:
        png = out / f"{int(num):03d}_{re.sub(r'[^0-9A-Za-zА-Яа-яЁё._-]', '_', str(name))}.png"
        try:
            try:
                ws = wb.Worksheets(sheet)
            except Exception:
                ws = wb.Worksheets(int(num) + 1)  # запасной вариант - по номеру листа
            ws.Activate()
            rng = ws.Range(addr)
        except Exception as e:  # noqa
            n_err += 1
            print(f"ERR  {sheet}: не удалось открыть лист/диапазон {addr}: {e}")
            continue
        why = capture_with_excel(ws, rng, png)
        if why is None:
            n_excel += 1
            print(f"OK   {png.name}")
            continue
        try:
            capture_by_drawing(ws, rng, png)
            n_draw += 1
            print(f"OK*  {png.name}  (нарисовано скриптом; Excel не отдал картинку: {why})")
        except Exception as e:  # noqa
            n_err += 1
            print(f"ERR  {sheet}: {why}; отрисовка: {e}")

    wb.Close(SaveChanges=False)
    excel.Quit()
    print(f"\nГотово. Папка: {out}\nснято Excel: {n_excel}, нарисовано скриптом: {n_draw}, ошибок: {n_err}")


if __name__ == "__main__":
    main()
