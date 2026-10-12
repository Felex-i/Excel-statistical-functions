#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Собирает Word-документ с одной большой таблицей по всем функциям:
№ | Функция | Ссылка | Синтаксис | Параметры | Скриншот

Данные берутся из excel_statistical_functions.xlsx (его делает build_excel.py),
скриншоты - из папки screenshots/ (их делает make_screenshots.py).
Скриншот подбирается по имени: 001_СРЗНАЧ.png  <->  функция №1 «СРЗНАЧ».

Нужен пакет:  pip install python-docx openpyxl pillow
Запуск:       python make_word.py
Свои пути:    python make_word.py файл.xlsx папка_со_скриншотами результат.docx
"""
import pathlib
import re
import sys

import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

XLSX = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "excel_statistical_functions.xlsx")
SHOTS = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else "screenshots")
OUT = pathlib.Path(sys.argv[3] if len(sys.argv) > 3 else "statistical_functions.docx")

# ширины колонок, см (сумма = 26,7 - вся ширина альбомной страницы A4 при полях 1,5 см)
WIDTHS = [1.0, 3.0, 3.8, 4.4, 6.5, 8.0]
HEADERS = ["№", "Функция", "Ссылка", "Синтаксис", "Параметры", "Скриншот"]


def shot_name(num, name):
    # то же правило, что в make_screenshots.py - так имена гарантированно совпадают
    return f"{int(num):03d}_{re.sub(r'[^0-9A-Za-zА-Яа-яЁё._-]', '_', str(name))}.png"


def find_shot(num, name):
    p = SHOTS / shot_name(num, name)
    if p.exists():
        return p
    found = sorted(SHOTS.glob(f"{int(num):03d}_*.png"))  # запасной вариант: по номеру
    return found[0] if found else None


def set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW")
        tcPr.append(w)
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_hyperlink(paragraph, url, text):
    rid = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.extend([color, u, sz])
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def fill_text(cell, text, bold=False, size=9):
    """Записывает текст в ячейку; каждая строка - отдельный абзац."""
    lines = [l for l in str(text or "").split("\n")] or [""]
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.bold = bold


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def main():
    if not XLSX.exists():
        sys.exit(f"Не найден файл {XLSX}. Сначала запустите build_excel.py")
    ws = openpyxl.load_workbook(XLSX, read_only=True)["Все функции"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1]]
    # колонки: 0 №, 1 Функция, 2 Англ., 3 Описание, 4 Синтаксис, 5 Аргументы, 6 Замечания, 7 Ссылка

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(9)

    doc.add_heading("Статистические функции Excel", level=1)
    doc.add_paragraph("Источник: support.microsoft.com/ru-ru/excel/statistical-functions-reference")

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")  # фиксированные ширины колонок
    tblPr.append(lay)
    for col, w in zip(table.columns, WIDTHS):
        col.width = Cm(w)
    for i, (h, w) in enumerate(zip(HEADERS, WIDTHS)):
        c = table.rows[0].cells[i]
        set_cell_width(c, w)
        shade(c, "D9E2F3")
        fill_text(c, h, bold=True, size=10)
    repeat_header(table.rows[0])

    missing = []
    for r in rows:
        num, name, _en, _desc, syntax, args, _rem, url = r[:8]
        row = table.add_row()
        no_split(row)
        cells = row.cells
        for i, w in enumerate(WIDTHS):
            set_cell_width(cells[i], w)
        fill_text(cells[0], num)
        fill_text(cells[1], name, bold=True, size=10)
        p = cells[2].paragraphs[0]
        if url:
            add_hyperlink(p, url, str(url).replace("https://", ""))
        fill_text(cells[3], syntax)
        fill_text(cells[4], args)
        shot = find_shot(num, name)
        p = cells[5].paragraphs[0]
        if shot:
            with Image.open(shot) as im:
                w_px, h_px = im.size
            width_cm = WIDTHS[5] - 0.3
            height_cm = width_cm * h_px / w_px
            if height_cm > 11:  # очень высокая картинка - ограничиваем по высоте
                width_cm *= 11 / height_cm
            p.add_run().add_picture(str(shot), width=Cm(width_cm))
            cap = cells[5].add_paragraph()
            run = cap.add_run(shot.name)
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        else:
            fill_text(cells[5], "(нет скриншота)")
            missing.append(str(name))

    doc.save(OUT)
    print(f"Готово: {OUT.resolve()}  (функций: {len(rows)})")
    if missing:
        print("Нет скриншота для:", ", ".join(missing))


if __name__ == "__main__":
    main()
